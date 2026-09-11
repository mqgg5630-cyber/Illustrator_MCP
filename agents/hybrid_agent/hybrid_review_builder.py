#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
hybrid_review_builder.py
中英双轨混合学术智能体 (Hybrid Academic Agent) 学术专著/学位论文 DOCX 活体编译引擎
1. 完整克隆高校权威标准模板，保留校徽、图徽与版式元数据；
2. 严格遵循 6-Section 分节拓扑架构与页眉隔离铁律，彻底根除“页眉全是目录”；
3. 支持中英文献混合活体引用（中文文献“等”、英文文献“et al.”）；
4. 正文流式注入 ADDIN ZOTERO_ITEM CSL_CITATION 活体复合域；
5. 文末参考文献精准包裹 ADDIN ZOTERO_BIBL CSL_BIBLIOGRAPHY 活体复合域；
6. 底层 OpenXML 属性分段切片（<=250字符）写入 docProps/custom.xml，ZIP 容器绝对唯一，零修复弹窗。
"""

import os
import sys
import json
import re
import time
import zipfile
import shutil
import tempfile
from xml.sax.saxutils import escape

# 强制重定向临时目录至 E 盘，严格恪守零 C 盘占用铁律（可通过 HUB_SCRATCH_DIR 覆盖；目录不可用时不阻断 import）
SCRATCH_DIR = os.environ.get("HUB_SCRATCH_DIR", r"E:\0mcp-agv\scratch")
try:
    if not os.path.isabs(SCRATCH_DIR):
        raise OSError("scratch dir must be absolute on this platform")
    os.makedirs(SCRATCH_DIR, exist_ok=True)
    os.environ["TEMP"] = SCRATCH_DIR
    os.environ["TMP"] = SCRATCH_DIR
    tempfile.tempdir = SCRATCH_DIR
except OSError:
    pass

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

TEMPLATE_PATH = os.environ.get("HUB_THESIS_TEMPLATE", r"E:\0writing\Lark-Formatter\鲁东大学学术学位论文_Zotero活动引用版_new.docx")
STYLE_GB7714 = "http://www.zotero.org/styles/china-national-standard-gb-t-7714-2015-numeric"

def set_cell_border(cell, **kwargs):
    """设置三线表标准科技边框"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{kwargs.get("top", "none")}" w:sz="{kwargs.get("top_sz", "4")}" w:space="0" w:color="{kwargs.get("top_color", "auto")}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:bottom w:val="{kwargs.get("bottom", "none")}" w:sz="{kwargs.get("bottom_sz", "4")}" w:space="0" w:color="{kwargs.get("bottom_color", "auto")}"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

class HybridReviewBuilder:
    def __init__(self, theme_dir):
        self.theme_dir = theme_dir
        self.manifest_file = os.path.join(theme_dir, "manifest.json")
        self.output_docx = os.path.join(theme_dir, f"{os.path.basename(theme_dir)}_学术专著论文.docx")
        
        with open(self.manifest_file, "r", encoding="utf-8") as f:
            self.manifest = json.load(f)
            
        self.item_map = {it["zotero_key"]: it for it in self.manifest if "zotero_key" in it}
        self.cite_counter = 0
        self.outline_file = os.path.join(theme_dir, "outline.json")
        self.outline = {}
        self.cite_order = []
        self.cite_num = {}

    def add_clean_heading(self, doc, text, level):
        """添加自动剥离显式序号的纯文本标题，100% 免疫双重标题"""
        clean_text = text
        clean_text = re.sub(r"^第\s*[0-9一二三四五六七八九十]+\s*章\s*", "", clean_text)
        clean_text = re.sub(r"^Chapter\s*[0-9]+\s*", "", clean_text, flags=re.IGNORECASE)
        clean_text = re.sub(r"^[0-9]+(\.[0-9]+)*\s*", "", clean_text)
        clean_text = clean_text.strip()
        h = doc.add_heading(clean_text, level=level)
        h.paragraph_format.keep_with_next = True
        return h

    def add_zotero_citation(self, paragraph, keys: list, display_text: str):
        """流式注入标准 ADDIN ZOTERO_ITEM CSL_CITATION 活体复合域"""
        self.cite_counter += 1
        citation_items = []
        for k in keys:
            it = self.item_map.get(k, {})
            title = it.get("title", "Article")
            journal = it.get("journal", "Academic Journal")
            year = str(it.get("year", 2024))
            authors_list = it.get("authors", ["Research Group"])
            lang = it.get("lang", "en")

            c_authors = []
            for a in authors_list:
                if lang == "zh":
                    c_authors.append({"family": a.strip(), "given": ""})
                else:
                    parts = a.strip().split()
                    if len(parts) > 1:
                        c_authors.append({"family": parts[-1], "given": " ".join(parts[:-1])})
                    else:
                        c_authors.append({"family": a.strip(), "given": ""})
            if not c_authors:
                c_authors = [{"family": "Author", "given": ""}]

            citation_items.append({
                "id": k,
                "uris": [f"http://zotero.org/users/local/user/items/{k}"],
                "itemData": {
                    "id": k,
                    "type": "article-journal",
                    "title": title,
                    "container-title": journal,
                    "issued": {"date-parts": [[int(year) if year.isdigit() else 2024]]},
                    "author": c_authors,
                    **({"DOI": it["doi"]} if it.get("doi") else {}),
                    **({"volume": str(it["volume"])} if it.get("volume") else {}),
                    **({"issue": str(it["issue"])} if it.get("issue") else {}),
                    **({"page": str(it["pages"])} if it.get("pages") else {}),
                    **({"language": "zh-CN"} if lang == "zh" else {}),
                }
            })

        payload = {
            "citationID": f"cHYB_{self.cite_counter:03d}",
            "properties": {
                "formattedCitation": display_text,
                "plainCitation": display_text,
                "noteIndex": 0,
                "dontUpdate": False
            },
            "citationItems": citation_items,
            "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"
        }

        instr_text = f' ADDIN ZOTERO_ITEM CSL_CITATION {json.dumps(payload, ensure_ascii=False, separators=(",", ":"))} '

        run_begin = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % nsdecls('w'))
        paragraph._p.append(run_begin)

        run_instr = parse_xml(r'<w:r %s><w:instrText xml:space="preserve">%s</w:instrText></w:r>' % (nsdecls('w'), escape(instr_text)))
        paragraph._p.append(run_instr)

        run_sep = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % nsdecls('w'))
        paragraph._p.append(run_sep)

        # 可见上标渲染 (高亮克莱因蓝)
        run_disp = paragraph.add_run(display_text)
        run_disp.font.superscript = True
        run_disp.font.size = Pt(10.5)
        run_disp.font.name = "Times New Roman"
        run_disp.font.color.rgb = RGBColor(0, 47, 167)

        run_end = parse_xml(r'<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % nsdecls('w'))
        paragraph._p.append(run_end)

    def add_para_with_cites(self, doc, text_segments):
        """向正文段落中流式写入文本与 Zotero 引注"""
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Inches(0.35)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for seg in text_segments:
            if isinstance(seg, str):
                seg = (seg,)
            if len(seg) == 1:
                t = seg[0]
                run = p.add_run(t)
                run.font.name = "宋体"
                run.font.size = Pt(11)
            elif len(seg) == 3:
                t, c_keys, disp = seg
                if t:
                    run = p.add_run(t)
                    run.font.name = "宋体"
                    run.font.size = Pt(11)
                keys_list = [c_keys] if isinstance(c_keys, str) else list(c_keys)
                self.add_zotero_citation(p, keys_list, disp)

    # ------------------------------------------------------------------
    # 大纲加载与引用编号解析
    # ------------------------------------------------------------------
    def load_outline(self):
        """加载课题目录下 outline.json（由 Antigravity 按真实文献撰写），并校验引用键。"""
        if not os.path.exists(self.outline_file):
            skel = write_outline_skeleton(self.theme_dir)
            raise FileNotFoundError(
                f"未找到大纲文件: {self.outline_file}\n"
                f"已生成骨架 {skel}，请让 Antigravity 通读文献后填写并另存为 outline.json；"
                f"完整示例参见 agents/hybrid_agent/examples/outline.collagen_example.json"
            )
        with open(self.outline_file, "r", encoding="utf-8") as f:
            outline = json.load(f)

        # 引用编号：按正文首次出现顺序编号 (GB/T 7714 顺序编码制)，
        # 参考文献列表按此顺序重排，保证 [n] 与文末条目严格一致。
        order = []
        unknown = set()
        def _register(keys):
            for k in keys:
                if k == "__ALL__":
                    for mk in self.item_map:
                        if mk not in order:
                            order.append(mk)
                    continue
                if k not in self.item_map:
                    unknown.add(k)
                    continue
                if k not in order:
                    order.append(k)
        for ch in outline.get("chapters", []):
            for sec in ch.get("sections", []):
                for blk in sec.get("blocks", []):
                    if blk.get("type", "paragraph") == "paragraph":
                        for seg in blk.get("segments", []):
                            if seg.get("cite"):
                                _register(seg["cite"])
        if unknown:
            raise ValueError(
                f"outline.json 中引用了 manifest.json 不存在的 zotero_key: {sorted(unknown)}\n"
                f"可用 key: {sorted(self.item_map)}"
            )
        # 未被正文引用的文献追加在末尾，仍进入参考文献并在总结段以 __ALL__ 覆盖
        for mk in self.item_map:
            if mk not in order:
                order.append(mk)
        self.cite_order = order
        self.cite_num = {k: i for i, k in enumerate(order, start=1)}
        self.outline = outline
        return outline

    def format_cite_label(self, keys):
        """将 zotero_key 列表转为 [1] / [2,5] / [1-10] 形式的上标文本。"""
        if keys == ["__ALL__"] or "__ALL__" in keys:
            nums = sorted(self.cite_num.values())
        else:
            nums = sorted({self.cite_num[k] for k in keys})
        if not nums:
            return ""
        # 连续区间压缩
        parts, start, prev = [], nums[0], nums[0]
        for n in nums[1:] + [None]:
            if n is not None and n == prev + 1:
                prev = n
                continue
            parts.append(f"{start}-{prev}" if prev - start >= 2 else (f"{start},{prev}" if prev != start else f"{start}"))
            if n is not None:
                start = prev = n
        return "[" + ",".join(parts) + "]"

    def resolve_keys(self, keys):
        if "__ALL__" in keys:
            return list(self.cite_order)
        return list(keys)

    # ------------------------------------------------------------------
    # 模板前置部分：封面 / 摘要 / 目录
    # ------------------------------------------------------------------
    @staticmethod
    def _set_run(paragraph, idx, text):
        if idx < len(paragraph.runs):
            paragraph.runs[idx].text = text

    def update_front_matter_in_place(self, doc):
        """原地更新封面、扉页与独创性声明（内容取自 outline.meta）"""
        print("  📝 原地更新外封面、中英文扉页元数据与原创性声明...")
        m = self.outline.get("meta", {})
        g = lambda k, d="": str(m.get(k, d))
        P = doc.paragraphs

        # 封面标题
        self._set_run(P[4], 0, g("title_zh_line1"))
        self._set_run(P[4], 1, "\n")
        self._set_run(P[4], 2, g("title_zh_line2"))

        # 封面 6行2列信息表
        table0 = doc.tables[0]
        for r, key in enumerate(["author_zh", "supervisor_zh", "major_zh", "direction_zh", "school_zh", "date_zh"]):
            cell_p = table0.cell(r, 1).paragraphs[0]
            if cell_p.runs:
                cell_p.runs[0].text = g(key)
            else:
                cell_p.add_run(g(key))

        # 中文扉页
        self._set_run(P[9], 0, g("title_zh_line1"))
        self._set_run(P[9], 1, "\n")
        self._set_run(P[9], 2, g("title_zh_line2"))
        self._set_run(P[11], 0, f"作者姓名：{g('author_zh')}")
        self._set_run(P[11], 2, f"指导教师：{g('supervisor_zh')}")
        self._set_run(P[11], 4, f"学科专业：{g('major_full_zh', g('major_zh'))}")
        self._set_run(P[11], 6, f"研究方向：{g('direction_zh')}")
        self._set_run(P[11], 10, g("school_full_zh", g("school_zh")))
        self._set_run(P[11], 12, g("date_zh_cn", g("date_zh")))

        # 英文扉页
        self._set_run(P[14], 0, g("title_en") + "\n\n\n")
        self._set_run(P[15], 0, f"{g('degree_en', 'M.D. Candidate')}: {g('author_en')}")
        self._set_run(P[15], 2, f"Supervisor: {g('supervisor_en')}")
        self._set_run(P[15], 4, f"Major: {g('major_en')}")
        self._set_run(P[15], 6, f"Research Interests: {g('direction_en')}")
        self._set_run(P[15], 10, g("school_en"))
        self._set_run(P[15], 12, g("date_en"))

        # 声明页签名
        sig = f"作者签名：{g('author_zh')}                                  日期：{g('date_zh')}"
        self._set_run(P[20], 0, sig)
        self._set_run(P[24], 0, sig)
        sup_first = g("supervisor_zh").split("/")[0].strip()
        self._set_run(P[24], 2, f"导师签名：{sup_first}                                  日期：{g('date_zh')}")

    def update_abstracts_in_place(self, doc):
        """原地更新中英文摘要与关键词（内容取自 outline）"""
        print("  📝 原地更新中英文摘要与学术关键词...")
        P = doc.paragraphs
        o = self.outline

        def _join(v):
            return "\n".join(v) if isinstance(v, list) else str(v)

        zh = o.get("abstract_zh", [])
        zh = zh if isinstance(zh, list) else [zh]
        en = o.get("abstract_en", [])
        en = en if isinstance(en, list) else [en]

        P[26].text = "摘  要"; P[26].style = "Front Matter Heading Unnumbered"
        P[27].text = zh[0] if zh else ""; P[27].style = "Normal"
        P[28].text = "\n".join(zh[1:]); P[28].style = "Normal"
        P[29].text = f"关键词：{o.get('keywords_zh', '')}"; P[29].style = "Normal"

        P[30].text = "Abstract"; P[30].style = "Front Matter Heading Unnumbered"
        P[31].text = en[0] if en else ""; P[31].style = "Normal"
        P[32].text = "\n".join(en[1:]); P[32].style = "Normal"
        P[33].text = f"KeyWords: {o.get('keywords_en', '')}"; P[33].style = "Normal"

    def update_toc_in_place(self, doc):
        """原地更新目录：条目由 outline.chapters 自动生成；页码由 Word 更新域时刷新，此处写占位页码。"""
        print("  📝 原地更新目录体系...")
        p34 = doc.paragraphs[34]
        p34.text = "目  录"
        p34.style = "TOC Heading"

        toc_data = [(1, "摘  要", "I"), (1, "Abstract", "II")]
        for ci, ch in enumerate(self.outline.get("chapters", []), start=1):
            toc_data.append((1, self._chapter_title(ci, ch), ""))
            for si, sec in enumerate(ch.get("sections", []), start=1):
                toc_data.append((2, self._section_title(ci, si, sec), ""))
        toc_data += [(1, "参考文献", ""), (1, "致谢", "")]

        first_body_idx = 35
        last_body_idx = 61  # 模板中正文起点前一段
        slots = last_body_idx - first_body_idx
        if len(toc_data) > slots:
            print(f"  ⚠️ 目录条目 {len(toc_data)} 条超过模板可容纳 {slots} 条，超出部分将被省略（请在 Word 中更新目录域）")
            toc_data = toc_data[:slots]

        for i, (level, title, page) in enumerate(toc_data):
            idx = first_body_idx + i
            if idx < len(doc.paragraphs):
                p = doc.paragraphs[idx]
                p.text = f"{title}\t{page}" if page else title
                p.style = f"TOC {level}"

        for j in range(first_body_idx + len(toc_data), last_body_idx):
            if j < len(doc.paragraphs):
                doc.paragraphs[j].text = ""
                doc.paragraphs[j].style = "Normal"

    # ------------------------------------------------------------------
    # 标题编号：大纲里写不写“第N章 / N.M”都可以，统一由此规范化
    # ------------------------------------------------------------------
    @staticmethod
    def _strip_num(text):
        return re.sub(r"^\s*(第\s*[0-9一二三四五六七八九十]+\s*章|\d+(?:\.\d+)*)\s*[、．.\s]*", "", text).strip()

    def _chapter_title(self, ci, ch):
        return f"第{ci}章 {self._strip_num(ch['title'])}"

    def _section_title(self, ci, si, sec):
        return f"{ci}.{si} {self._strip_num(sec['title'])}"

    # ------------------------------------------------------------------
    # 正文块渲染
    # ------------------------------------------------------------------
    def render_paragraph_block(self, doc, blk):
        segs = []
        for seg in blk.get("segments", []):
            text = seg.get("text", "")
            keys = seg.get("cite")
            if keys:
                keys = [keys] if isinstance(keys, str) else list(keys)
                label = seg.get("label") or self.format_cite_label(keys)
                segs.append((text, self.resolve_keys(keys), label))
            else:
                segs.append((text,))
        self.add_para_with_cites(doc, segs)

    def render_table_block(self, doc, blk, ci, table_no):
        caption = blk.get("caption", "")
        caption = re.sub(r"^\s*表\s*[\d\-–]+\s*", "", caption)
        p_tbl_title = doc.add_paragraph(f"表{ci}-{table_no} {caption}")
        p_tbl_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tbl_title.runs[0].font.bold = True
        p_tbl_title.runs[0].font.name = "宋体"
        p_tbl_title.runs[0].font.size = Pt(10.5)

        headers = blk.get("headers", [])
        rows_data = blk.get("rows", [])
        if not headers or not rows_data:
            return
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for col_idx, h in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = str(h)
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.name = "宋体"
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_border(cell, top="single", top_sz="12", bottom="single", bottom_sz="6")

        for r_idx, rdata in enumerate(rows_data, start=1):
            for col_idx in range(len(headers)):
                val = str(rdata[col_idx]) if col_idx < len(rdata) else ""
                cell = table.cell(r_idx, col_idx)
                cell.text = val
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].font.name = "宋体" if any('\u4e00' <= c <= '\u9fff' for c in val) else "Times New Roman"
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                bottom_sz = "12" if r_idx == len(rows_data) else "0"
                bottom_val = "single" if r_idx == len(rows_data) else "none"
                set_cell_border(cell, bottom=bottom_val, bottom_sz=bottom_sz)

        # 行跨页防撕裂 (cantSplit) 与跨页表头重复 (tblHeader)；CT_TrPr 要求 cantSplit 在 tblHeader 之前
        for row in table.rows:
            row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        table.rows[0]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)

    # ------------------------------------------------------------------
    # 参考文献条目格式化 (GB/T 7714-2015 顺序编码制，Zotero 刷新后会被 CSL 覆盖)
    # ------------------------------------------------------------------
    @staticmethod
    def format_reference(idx, it):
        lang = it.get("lang", "en")
        authors_list = it.get("authors") or (["佚名"] if lang == "zh" else ["Anon"])
        if lang == "zh":
            authors_str = "，".join(authors_list[:3]) + ("，等" if len(authors_list) > 3 else "")
        else:
            authors_str = ", ".join(authors_list[:3]) + (", et al" if len(authors_list) > 3 else "")
        title = it.get("title", "")
        journal = it.get("journal", "")
        year = it.get("year", "")
        vol = it.get("volume", "")
        issue = it.get("issue", "")
        pages = it.get("pages", "")
        doi = it.get("doi", "")
        vi = f"{vol}" + (f"({issue})" if issue else "") if vol else ""
        tail = ", ".join(x for x in [str(year), vi] if x)
        if pages:
            tail += f": {pages}"
        doi_str = f" DOI: {doi}." if doi else ""
        return f"[{idx}] {authors_str}. {title}[J]. {journal}, {tail}.{doi_str}"

    # ------------------------------------------------------------------
    # 主编译入口
    # ------------------------------------------------------------------
    def build_docx(self):
        print("=" * 65)
        print(f"🚀 [Hybrid Review Builder] 启动中英双轨学术专著活体编译: {self.theme_dir}")
        print(f"📚 可用真实中英文献篇数: {len(self.manifest)}")
        print("=" * 65)

        outline = self.load_outline()
        print(f"📑 大纲章节: {len(outline.get('chapters', []))} 章，文献引用顺序已解析 ({len(self.cite_order)} 篇)")

        doc = docx.Document(TEMPLATE_PATH)

        # 1~3. 封面 / 摘要 / 目录
        self.update_front_matter_in_place(doc)
        self.update_abstracts_in_place(doc)
        self.update_toc_in_place(doc)

        # 4. 彻底清空模板原有旧正文（保留 Section 3 之前的完整格式）
        body = doc._body._body
        start_p_elem = doc.paragraphs[61]._element if len(doc.paragraphs) > 61 else None
        if start_p_elem is not None:
            children = list(body)
            if start_p_elem in children:
                start_idx = children.index(start_p_elem)
                for ch_elem in children[start_idx:]:
                    body.remove(ch_elem)

        # 5. 正文章节（全部来自 outline.chapters）
        for ci, ch in enumerate(outline.get("chapters", []), start=1):
            self.add_clean_heading(doc, self._chapter_title(ci, ch), level=1)
            table_no = 0
            for si, sec in enumerate(ch.get("sections", []), start=1):
                self.add_clean_heading(doc, self._section_title(ci, si, sec), level=2)
                for blk in sec.get("blocks", []):
                    btype = blk.get("type", "paragraph")
                    if btype == "paragraph":
                        self.render_paragraph_block(doc, blk)
                    elif btype == "table":
                        table_no += 1
                        self.render_table_block(doc, blk, ci, table_no)
                    else:
                        print(f"  ⚠️ 未知块类型 {btype}，已跳过")

        # 6. Section 4 结束，注入标准分节符 s4_xml
        p_space_end = doc.add_paragraph()
        s4_xml = (
            f'<w:sectPr {nsdecls("w")} {nsdecls("r")}>\n'
            f'  <w:headerReference r:id="rId13" w:type="default"/>\n'
            f'  <w:footerReference r:id="rId14" w:type="default"/>\n'
            f'  <w:pgSz w:w="11906" w:h="16838"/>\n'
            f'  <w:pgMar w:top="1701" w:right="1701" w:bottom="1417" w:left="1701" w:header="1134" w:footer="1247" w:gutter="0"/>\n'
            f'  <w:pgNumType w:fmt="decimal" w:start="1"/>\n'
            f'  <w:cols w:space="720" w:num="1"/>\n'
            f'  <w:docGrid w:linePitch="360" w:charSpace="0"/>\n'
            f'</w:sectPr>'
        )
        p_space_end._p.get_or_add_pPr().append(parse_xml(s4_xml))

        # 7. Section 5: 参考文献（按正文首次引用顺序）与致谢
        p_ref_h = doc.add_paragraph("参 考 文 献")
        p_ref_h.style = "Heading 1 Unnumbered"
        p_ref_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ref_h.paragraph_format.space_before = Pt(16)
        p_ref_h.paragraph_format.space_after = Pt(12)

        for k in self.cite_order:
            it = self.item_map[k]
            idx = self.cite_num[k]
            lang = it.get("lang", "en")
            p_ref = doc.add_paragraph(self.format_reference(idx, it))
            p_ref.paragraph_format.line_spacing = 1.15
            p_ref.paragraph_format.space_after = Pt(3)
            p_ref.paragraph_format.first_line_indent = Inches(0)
            for r in p_ref.runs:
                r.font.name = "宋体" if lang == "zh" else "Times New Roman"
                r.font.size = Pt(10)

        doc.add_page_break()
        p_ack = doc.add_paragraph("致  谢")
        p_ack.style = "Heading 1 Unnumbered"
        p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ack.paragraph_format.space_before = Pt(16)
        p_ack.paragraph_format.space_after = Pt(12)
        ack = outline.get("acknowledgement", [])
        ack = ack if isinstance(ack, list) else [ack]
        for para in ack:
            p_ack_text = doc.add_paragraph(para)
            p_ack_text.paragraph_format.line_spacing = 1.25
            p_ack_text.paragraph_format.first_line_indent = Inches(0.35)
            for r in p_ack_text.runs:
                r.font.name = "宋体"
                r.font.size = Pt(11)

        # Section 5 终节点属性 (绑定 header8.xml "Heading 1 Unnumbered")
        s5_xml = (
            f'<w:sectPr {nsdecls("w")} {nsdecls("r")}>\n'
            f'  <w:headerReference r:id="rId19" w:type="default"/>\n'
            f'  <w:footerReference r:id="rId20" w:type="default"/>\n'
            f'  <w:pgSz w:w="11906" w:h="16838"/>\n'
            f'  <w:pgMar w:top="1701" w:right="1701" w:bottom="1417" w:left="1701" w:header="1134" w:footer="1247" w:gutter="0"/>\n'
            f'  <w:pgNumType w:fmt="decimal"/>\n'
            f'  <w:cols w:space="720" w:num="1"/>\n'
            f'  <w:docGrid w:linePitch="360" w:charSpace="0"/>\n'
            f'</w:sectPr>'
        )
        body.append(parse_xml(s5_xml))

        doc.save(self.output_docx)
        print(f"  💾 基础 DOCX 结构已保存至: {self.output_docx}")

        # OpenXML 复合域后处理与页眉净化
        self.wrap_bibliography_field(self.output_docx)
        self.patch_zotero_preferences(self.output_docx)
        self.patch_headers_fallback(self.output_docx)
        self.verify_docx(self.output_docx)
        return self.output_docx

    def wrap_bibliography_field(self, docx_path: str):
        """将 参 考 文 献 至 致  谢 之间的列表包裹在 ADDIN ZOTERO_BIBL 复合域内"""
        tmp_path = docx_path + ".tmp.docx"
        with zipfile.ZipFile(docx_path, 'r') as zin, zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/document.xml":
                    xml_str = data.decode('utf-8')
                    ref_pos = xml_str.find("参 考 文 献")
                    ack_pos = xml_str.find("致  谢")
                    if ref_pos >= 0 and ack_pos > ref_pos:
                        bib_slice = xml_str[ref_pos:ack_pos]
                        ref_p_pattern = re.compile(r'(<w:p\b[^>]*>(?:(?!</w:p>).)*?<w:t\b[^>]*>\[\d+\]\s+.*?</w:p>)', re.DOTALL)
                        matches = list(ref_p_pattern.finditer(bib_slice))
                        if matches:
                            first_m = matches[0]
                            last_m = matches[-1]
                            bib_start = (
                                '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                                '<w:r><w:instrText xml:space="preserve"> ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY </w:instrText></w:r>'
                                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                            )
                            bib_end = '<w:r><w:fldChar w:fldCharType="end"/></w:r>'

                            first_p = first_m.group(1)
                            last_p = last_m.group(1)

                            p_pos = first_p.find("</w:pPr>")
                            if p_pos >= 0:
                                p_pos += len("</w:pPr>")
                                new_first_p = first_p[:p_pos] + bib_start + first_p[p_pos:]
                            else:
                                new_first_p = "<w:p>" + bib_start + first_p[len("<w:p>"):]

                            end_pos = last_p.rfind("</w:p>")
                            new_last_p = last_p[:end_pos] + bib_end + last_p[end_pos:]

                            if len(matches) == 1:
                                comb = first_p
                                p_pos = comb.find("</w:pPr>") + len("</w:pPr>")
                                comb = comb[:p_pos] + bib_start + comb[p_pos:]
                                end_pos = comb.rfind("</w:p>")
                                comb = comb[:end_pos] + bib_end + comb[end_pos:]
                                new_slice = bib_slice[:first_m.start()] + comb + bib_slice[first_m.end():]
                            else:
                                new_slice = bib_slice[:first_m.start()] + new_first_p + bib_slice[first_m.end():last_m.start()] + new_last_p + bib_slice[last_m.end():]

                            xml_str = xml_str[:ref_pos] + new_slice + xml_str[ack_pos:]
                            data = xml_str.encode('utf-8')
                            print("  ✅ [Zotero Live] 成功将中英双轨参考文献包裹为 ADDIN ZOTERO_BIBL 复合域！")
                zout.writestr(item, data)

        if os.path.exists(docx_path):
            os.remove(docx_path)
        os.rename(tmp_path, docx_path)

    def patch_zotero_preferences(self, docx_path: str):
        """向 docProps/custom.xml 注册 <=250 字符分段切片与 Zotero 活体首选项"""
        csl_prefs = {
            "style": {"styleID": STYLE_GB7714, "locale": "zh-CN", "hasBibliography": True},
            "prefs": {"displayAs": "numeric", "notesType": "numeric", "automaticJournalAbbreviations": False}
        }
        prefs_json = json.dumps(csl_prefs, ensure_ascii=False)
        chunk_size = 250
        chunks = [prefs_json[i:i+chunk_size] for i in range(0, len(prefs_json), chunk_size)]

        custom_xml_lines = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        ]
        pid = 2
        for idx, ch in enumerate(chunks, start=1):
            custom_xml_lines.append(f'  <property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{pid}" name="ZOTERO_PREF_{idx}"><vt:lpwstr>{escape(ch)}</vt:lpwstr></property>')
            pid += 1
        custom_xml_lines.append('</Properties>')
        custom_xml_content = "\n".join(custom_xml_lines).encode('utf-8')

        tmp_path = docx_path + ".tmp.docx"
        with zipfile.ZipFile(docx_path, 'r') as zin, zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
            has_custom = False
            for item in zin.infolist():
                if item.filename == "docProps/custom.xml":
                    continue
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    xml_str = data.decode('utf-8')
                    if 'custom-properties' not in xml_str:
                        xml_str = xml_str.replace('</Types>', '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
                        data = xml_str.encode('utf-8')
                elif item.filename == "_rels/.rels":
                    xml_str = data.decode('utf-8')
                    if 'custom-properties' not in xml_str:
                        xml_str = xml_str.replace('</Relationships>', '<Relationship Id="rIdCustomProps" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
                        data = xml_str.encode('utf-8')
                zout.writestr(item, data)
            zout.writestr("docProps/custom.xml", custom_xml_content)

        if os.path.exists(docx_path):
            os.remove(docx_path)
        os.rename(tmp_path, docx_path)
        print("  ✅ [Zotero Live] 成功写入 docProps/custom.xml 250字符分段切片与 Zotero 活体首选项！")

    def patch_headers_fallback(self, docx_path: str):
        """净化页眉缓存回退文本，防止 Word 在未自动刷新域时错误显示'目录'或中文回退"""
        chapters = self.outline.get("chapters", [])
        first_ch_title = self._strip_num(chapters[0]["title"]) if chapters else "绪论"
        tmp_path = docx_path + ".tmp.docx"
        with zipfile.ZipFile(docx_path, 'r') as zin, zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == 'word/header4.xml':
                    xml_str = data.decode('utf-8')
                    xml_str = re.sub(r'<w:t>[^<]*?</w:t>', '<w:t>目  录</w:t>', xml_str)
                    data = xml_str.encode('utf-8')
                elif item.filename in ['word/header5.xml', 'word/header6.xml', 'word/header7.xml']:
                    xml_str = data.decode('utf-8')
                    xml_str = re.sub(r'<w:t>[^<]*?第1章[^<]*?</w:t>', '<w:t>第1章　</w:t>', xml_str)
                    xml_str = re.sub(r'<w:t>[^<]*?引言[^<]*?</w:t>', f'<w:t>{escape(first_ch_title)}</w:t>', xml_str)
                    data = xml_str.encode('utf-8')
                elif item.filename in ['word/header8.xml', 'word/header9.xml']:
                    xml_str = data.decode('utf-8')
                    xml_str = re.sub(r'<w:t>[^<]*?参考文献[^<]*?</w:t>', '<w:t>参 考 文 献</w:t>', xml_str)
                    data = xml_str.encode('utf-8')
                zout.writestr(item, data)

        if os.path.exists(docx_path):
            os.remove(docx_path)
        os.rename(tmp_path, docx_path)
        print("  ✨ 页眉回退文本与 6-Section 多节架构净化完成！")

    # ------------------------------------------------------------------
    # 编译后自检
    # ------------------------------------------------------------------
    def verify_docx(self, docx_path: str):
        """检查 6-Section 分节数、活体引注域数量、参考文献域包裹与 ZIP 条目唯一性。"""
        problems = []
        with zipfile.ZipFile(docx_path, "r") as z:
            names = z.namelist()
            dup = {n for n in names if names.count(n) > 1}
            if dup:
                problems.append(f"ZIP 重复条目: {sorted(dup)}")
            xml_str = z.read("word/document.xml").decode("utf-8")
        n_sect = xml_str.count("<w:sectPr")
        n_cite = xml_str.count("ADDIN ZOTERO_ITEM CSL_CITATION")
        n_bibl = xml_str.count("ADDIN ZOTERO_BIBL")
        expected_cites = sum(
            1 for ch in self.outline.get("chapters", []) for s in ch.get("sections", [])
            for b in s.get("blocks", []) if b.get("type", "paragraph") == "paragraph"
            for seg in b.get("segments", []) if seg.get("cite")
        )
        if n_sect != 6:
            problems.append(f"分节数应为 6，实际 {n_sect}")
        if n_cite != expected_cites:
            problems.append(f"活体引注域 {n_cite} 个，大纲期望 {expected_cites} 个")
        if n_bibl != 1:
            problems.append(f"ZOTERO_BIBL 域应为 1 个，实际 {n_bibl}")
        if problems:
            print("  ⚠️ [自检] 发现问题:")
            for p in problems:
                print(f"     - {p}")
        else:
            print(f"  ✅ [自检] 6 节 / {n_cite} 处活体引注 / 1 个参考文献域 / ZIP 唯一 — 全部通过")
        return problems


def write_outline_skeleton(theme_dir: str, force: bool = False) -> str:
    """依据 manifest.json 生成 outline.skeleton.json：列出全部可引用文献与空章节骨架，
    供 Antigravity 阅读文献后填写正文并另存为 outline.json。"""
    manifest_file = os.path.join(theme_dir, "manifest.json")
    out = os.path.join(theme_dir, "outline.skeleton.json")
    if os.path.exists(out) and not force:
        return out
    with open(manifest_file, "r", encoding="utf-8") as f:
        manifest = json.load(f)
    def _ok(it):
        st = (it.get("verification") or {}).get("status")
        return st in (None, "verified", "suspicious")   # 未跑 verify 的旧 manifest 也放行
    refs = [
        {
            "zotero_key": it.get("zotero_key"),
            "lang": it.get("lang", "en"),
            "type": it.get("type", ""),
            "verification": (it.get("verification") or {}).get("status", "unchecked"),
            "title": it.get("title", ""),
            "authors": (it.get("authors") or [])[:3],
            "journal": it.get("journal", ""),
            "year": it.get("year", ""),
            "tldr": it.get("tldr", ""),
            "digest_card": f"digest/{it.get('zotero_key')}.md",
        }
        for it in manifest if it.get("zotero_key") and _ok(it)
    ]
    skeleton = {
        "_instructions": (
            "由 Antigravity 按 .agents/skills/review-writing/SKILL.md 填写：1) 先读 digest/INDEX.md 与 digest/literature_matrix.md，再按需打开每篇 digest_card（不读 PDF）；2) 填写 meta/摘要/关键词；"
            "3) 在 chapters 中撰写 4~6 章、每章 2~3 节；每节 blocks 为 paragraph 或 table；"
            "paragraph.segments 中每个 segment 为 {text, cite:[zotero_key,...]}，cite 只能填下方列出的 key；"
            "\"__ALL__\" 表示引用全部文献；正文中每个数值必须能在对应 digest_card 的 Quotable facts / Abstract / Results / Conclusion 中逐字找到；"
            "type=review 的文献只能用于背景/现状段；4) 删除本字段与 _available_references，另存为 outline.json。"
        ),
        "_available_references": refs,
        "meta": {k: "" for k in [
            "title_zh_line1", "title_zh_line2", "title_en", "author_zh", "author_en", "supervisor_zh", "supervisor_en",
            "major_zh", "major_full_zh", "major_en", "direction_zh", "direction_en", "school_zh", "school_full_zh",
            "school_en", "date_zh", "date_zh_cn", "date_en", "degree_en"]},
        "abstract_zh": ["", ""], "keywords_zh": "", "abstract_en": ["", ""], "keywords_en": "",
        "chapters": [
            {"title": "绪论", "sections": [{"title": "研究背景", "blocks": [
                {"type": "paragraph", "segments": [{"text": "……", "cite": [refs[0]["zotero_key"]] if refs else []}]}]}]},
            {"title": "总结与展望", "sections": [{"title": "主要结论", "blocks": [
                {"type": "paragraph", "segments": [{"text": "……", "cite": ["__ALL__"]}]}]}]},
        ],
        "acknowledgement": [""],
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(skeleton, f, ensure_ascii=False, indent=2)
    return out


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Hybrid Academic Agent — outline.json 驱动的学术专著 DOCX 编译器")
    ap.add_argument("theme_dir", help="课题目录（需含 manifest.json 与 outline.json）")
    ap.add_argument("--template", default=None, help="覆盖默认学位论文模板路径")
    ap.add_argument("--skeleton", action="store_true", help="仅依据 manifest.json 生成 outline.skeleton.json 供 Antigravity 填写")
    args = ap.parse_args()
    if args.template:
        TEMPLATE_PATH = args.template
    if args.skeleton:
        print("📝 已生成大纲骨架:", write_outline_skeleton(args.theme_dir, force=True))
        sys.exit(0)
    HybridReviewBuilder(args.theme_dir).build_docx()
