# -*- coding: utf-8 -*-
"""
全面测试 ARTA Agent 的多仓库 PPT 引擎集成与字号质检
"""
import sys
from pathlib import Path
import pptx
from pptx import Presentation

sys.path.insert(0, r"e:\0mcp-agv")
from agents.academic_thesis_agent.arta_agent import AcademicThesisAgent, ThesisStudentInfo

def inspect_pptx_font_sizes(pptx_file: str) -> dict:
    prs = Presentation(pptx_file)
    font_sizes = []
    text_count = 0
    below_18pt = []
    
    for s_idx, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    text_count += 1
                    # 检查段落字体或 run 字体
                    pt_size = None
                    if paragraph.font.size:
                        pt_size = paragraph.font.size.pt
                    for run in paragraph.runs:
                        if run.font.size:
                            pt_size = run.font.size.pt
                            break
                    
                    if pt_size is not None:
                        font_sizes.append(pt_size)
                        if pt_size < 18.0 and not any(tag in text.lower() for tag in ["page", "/", "consolas", "academic research", "2026"]):
                            below_18pt.append({
                                "slide": s_idx,
                                "text": text[:30],
                                "size": pt_size
                            })
                            
    min_size = min(font_sizes) if font_sizes else 0
    max_size = max(font_sizes) if font_sizes else 0
    return {
        "slide_count": len(prs.slides),
        "text_count": text_count,
        "min_size": min_size,
        "max_size": max_size,
        "below_18pt_count": len(below_18pt),
        "violations": below_18pt
    }

def main():
    agent = AcademicThesisAgent(workspace_dir=r"e:\0mcp-agv\ARTA_Agent_Output\PPT_Engine_Tests")
    
    engines = [
        ("ppt-master", "umami-peptide-ml"),
        ("cyber-ppt", None),
        ("swiss-pptx", None),
        ("guizang-ppt", None),
        ("banana-slides", None),
        ("dashi-ppt", "theme07")
    ]
    
    student = ThesisStudentInfo(
        school_name="鲁东大学",
        student_name="文  少",
        degree_field="工学 · 食品科学与工程",
        degree_type="硕士学位论文"
    )
    
    topic = "机器学习驱动的食源性鲜味肽高通量筛选与机制解析"
    
    print("="*70)
    print("  ARTA Agent PPT Engines Full Suite Test & 18pt Font Validation")
    print("="*70)
    
    summary_report = {}
    
    for engine_name, tpl in engines:
        print(f"\n[Testing Engine] --> {engine_name} (Template: {tpl})")
        res = agent.generate_presentation(
            topic=topic,
            ppt_engine=engine_name,
            ppt_template=tpl,
            student_info=student
        )
        pptx_path = res.get("pptx_file")
        if pptx_path and Path(pptx_path).exists():
            metrics = inspect_pptx_font_sizes(pptx_path)
            summary_report[engine_name] = {
                "status": "PASS",
                "pptx_file": pptx_path,
                "metrics": metrics
            }
            print(f"  [SUCCESS] Generated: {Path(pptx_path).name}")
            print(f"  [METRICS] Slides: {metrics['slide_count']} | Min Size: {metrics['min_size']}pt | Max Size: {metrics['max_size']}pt | Violations (<18pt content): {metrics['below_18pt_count']}")
        else:
            summary_report[engine_name] = {"status": "FAIL", "result": res}
            print(f"  [FAILED] No pptx file generated: {res}")
            
    print("\n" + "="*70)
    print("  ALL ENGINES TEST RESULTS SUMMARY")
    print("="*70)
    for k, v in summary_report.items():
        st = v["status"]
        if st == "PASS":
            m = v["metrics"]
            print(f"  [PASS] [{k:14s}] Slides: {m['slide_count']:02d} | Font Range: {m['min_size']}pt - {m['max_size']}pt | File: {Path(v['pptx_file']).name}")
        else:
            print(f"  [FAIL] [{k:14s}] FAILED")

if __name__ == "__main__":
    main()
