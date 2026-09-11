import docx
from docx.oxml import parse_xml
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

doc = docx.Document(r'E:\0writing\Lark-Formatter\鲁东大学学术学位论文_Zotero活动引用版_new.docx')
body = doc._body._body

# Remove old paragraphs from 61 onward
for pe in [p._p for p in doc.paragraphs[61:]]:
    if pe.getparent() is not None:
        pe.getparent().remove(pe)

# Remove old tables after table 0
for tbl in doc.tables[1:]:
    if tbl._tbl.getparent() is not None:
        tbl._tbl.getparent().remove(tbl._tbl)

s4_xml = (
    '<w:sectPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<w:headerReference r:id="rId13" w:type="default"/>'
    '<w:footerReference r:id="rId14" w:type="default"/>'
    '<w:pgSz w:w="11906" w:h="16838"/>'
    '<w:pgMar w:top="1701" w:right="1701" w:bottom="1417" w:left="1701" w:header="1134" w:footer="1247" w:gutter="0"/>'
    '<w:pgNumType w:fmt="decimal" w:start="1"/>'
    '<w:cols w:space="720" w:num="1"/>'
    '<w:docGrid w:linePitch="360" w:charSpace="0"/>'
    '</w:sectPr>'
)

old_body_sect = body.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr')
if old_body_sect is not None:
    body.remove(old_body_sect)

body.append(parse_xml(s4_xml))

print('New sections count:', len(doc.sections))
for i, s in enumerate(doc.sections):
    h_text = s.header.paragraphs[0].text if s.header.paragraphs else ''
    print(f'Section {i}: header={repr(h_text)}')
