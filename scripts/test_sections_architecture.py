import docx
from docx.oxml import parse_xml
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

doc = docx.Document(r'E:\0writing\Lark-Formatter\鲁东大学学术学位论文_Zotero活动引用版_new.docx')
body = doc._body._body

# Remove old paragraphs from 61 onward
for pe in [p._p for p in doc.paragraphs[61:]]:
    if pe.getparent() is not None:
        pe.getparent().remove(pe)

# Remove old tables after table 0
for tbl in doc.tables[1:]:
    if tbl._tbl.getparent() is not None:
        tbl._tbl.getparent().remove(tbl._tbl)

# Remove any leftover sectPr from body
old_body_sect = body.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr')
if old_body_sect is not None:
    body.remove(old_body_sect)

# Section 4 sectPr XML for Chapter body (ending at Chapter 5)
s4_xml = (
    '<w:sectPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<w:headerReference r:id="rId13" w:type="default"/>'
    '<w:footerReference r:id="rId14" w:type="default"/>'
    '<w:pgSz w:w="11906" w:h="16838"/>'
    '<w:pgMar w:top="1701" w:right="1701" w:bottom="1417" w:left="1701" w:header="1134" w:footer="1247" w:gutter="0"/>'
    '<w:pgNumType w:fmt="decimal" w:start="1"/>'
    '<w:cols w:space="720" w:num="1"/>'
    '<w:docGrid w:linePitch="360" w:charSpace="0"/>'
    '</w:sectPr>'
)

# Section 5 sectPr XML for References & Acknowledgements
s5_xml = (
    '<w:sectPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<w:headerReference r:id="rId19" w:type="default"/>'
    '<w:footerReference r:id="rId20" w:type="default"/>'
    '<w:pgSz w:w="11906" w:h="16838"/>'
    '<w:pgMar w:top="1701" w:right="1701" w:bottom="1417" w:left="1701" w:header="1134" w:footer="1247" w:gutter="0"/>'
    '<w:pgNumType w:fmt="decimal"/>'
    '<w:cols w:space="720" w:num="1"/>'
    '<w:docGrid w:linePitch="360" w:charSpace="0"/>'
    '</w:sectPr>'
)

# Add sample chapter paragraphs
p_ch1 = doc.add_paragraph("Introduction", style="Heading 1")
p_ch1_body = doc.add_paragraph("This is chapter 1 body text.")

p_ch5 = doc.add_paragraph("Perspectives", style="Heading 1")
p_ch5_body = doc.add_paragraph("This is chapter 5 body text.")

# The last paragraph of Section 4 (chapter body) receives s4_xml in its pPr!
pPr = p_ch5_body._p.get_or_add_pPr()
pPr.append(parse_xml(s4_xml))

# Now add References and Acknowledgements in Section 5
p_ref = doc.add_paragraph("References", style="Heading 1 Unnumbered")
p_ref_item = doc.add_paragraph("[1] Test Ref Item.")

p_ack = doc.add_paragraph("Acknowledgements", style="Heading 1 Unnumbered")
p_ack_text = doc.add_paragraph("Thank you all.")

# Finally append s5_xml at the end of body!
body.append(parse_xml(s5_xml))

print('Total sections:', len(doc.sections))
for i, s in enumerate(doc.sections):
    h = s.header.paragraphs[0].text if s.header.paragraphs else ''
    print(f'Section {i}: header={repr(h)}')
